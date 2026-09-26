"""
leads_import.py — Turns an uploaded file into leads.

Handles CSV and Excel. Expects the template's headers; anything else is
rejected with a clear message rather than guessed at, because guessing a
column wrong means calling the wrong number.

Phone numbers arrive in every imaginable format and are normalised to one,
so duplicate detection works and Vapi gets something it can dial.
"""

import io
import csv
import re
from dataclasses import dataclass, field
from typing import Iterator

from openpyxl import load_workbook

# The template's columns. Matched case-insensitively and ignoring
# surrounding whitespace, since spreadsheet exports add both.
REQUIRED_COLUMNS = ("name", "phone")
OPTIONAL_COLUMNS = ("language",)

VALID_LANGUAGES = {"english", "hindi", "telugu"}
DEFAULT_LANGUAGE = "english"

# Indian mobile numbers are ten digits starting 6, 7, 8 or 9. Landlines and
# short codes are not callable by this product, so they are rejected rather
# than silently accepted and failed at dial time.
MOBILE_PATTERN = re.compile(r"^[6-9]\d{9}$")

# Guard against someone uploading a million-row export
MAX_ROWS = 5000


@dataclass
class ParsedLead:
    row: int
    name: str
    phone: str          # normalised
    language: str


@dataclass
class RejectedRow:
    row: int
    reason: str
    raw: str = ""       # what was in the row, for the error report


@dataclass
class ParseResult:
    leads: list[ParsedLead] = field(default_factory=list)
    rejected: list[RejectedRow] = field(default_factory=list)

    @property
    def total(self) -> int:
        return len(self.leads) + len(self.rejected)


class ImportError_(Exception):
    """The file itself could not be read — wrong format, missing columns."""


# ── Phone normalisation ───────────────────────────────────────────────────────

def normalise_phone(raw: str) -> str | None:
    """
    Reduce any Indian mobile format to '+91 XXXXX XXXXX'.

    Accepts, among others:
        9876543210          +91 98765 43210     +919876543210
        091-9876543210      0 9876543210        91 9876543210

    Returns None if what remains is not a valid Indian mobile. Being strict
    here is deliberate — a malformed number that reaches Vapi fails at dial
    time, by which point nobody is watching.
    """
    if not raw:
        return None

    # Spreadsheets love turning phone numbers into floats
    text = str(raw).strip()
    if text.endswith(".0"):
        text = text[:-2]

    digits = re.sub(r"\D", "", text)

    # Strip the country code or trunk prefix, whichever is present
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    elif len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    elif len(digits) == 13 and digits.startswith("091"):
        digits = digits[3:]

    if not MOBILE_PATTERN.match(digits):
        return None

    return f"+91 {digits[:5]} {digits[5:]}"


def phone_digits(phone: str) -> str:
    """
    Just the digits, for comparison.

    Existing leads may be stored in a different spacing than newly imported
    ones, so duplicate checks compare on this rather than the display form.
    """
    return re.sub(r"\D", "", phone or "")[-10:]


# ── Reading files ─────────────────────────────────────────────────────────────

def _normalise_header(value) -> str:
    return str(value or "").strip().lower().replace(" ", "_")


def _check_headers(headers: list[str]) -> dict[str, int]:
    """
    Map each expected column to its position.

    Raises if a required column is missing, naming what was found — a user
    whose file fails needs to know which header to fix.
    """
    found = {h: i for i, h in enumerate(headers) if h}

    missing = [c for c in REQUIRED_COLUMNS if c not in found]
    if missing:
        raise ImportError_(
            f"Missing required column{'s' if len(missing) > 1 else ''}: "
            f"{', '.join(missing)}. "
            f"Found: {', '.join(headers) or 'nothing'}. "
            f"Download the template for the expected format."
        )

    return {c: found[c] for c in REQUIRED_COLUMNS + OPTIONAL_COLUMNS if c in found}


def _read_csv(content: bytes) -> Iterator[list]:
    # Excel on Windows writes UTF-8 with a BOM; utf-8-sig strips it so the
    # first header does not arrive as '\ufeffname'
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportError_("Could not read the file — unrecognised text encoding")

    for row in csv.reader(io.StringIO(text)):
        yield row


def _read_excel(content: bytes) -> Iterator[list]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportError_(f"Could not read the Excel file: {exc}") from exc

    sheet = workbook.active
    if sheet is None:
        raise ImportError_("The workbook has no sheets")

    for row in sheet.iter_rows(values_only=True):
        yield list(row)


def parse_file(filename: str, content: bytes) -> ParseResult:
    """
    Read an uploaded file into leads.

    Rows that cannot be used are collected rather than aborting the import —
    one bad phone number in row 40 should not discard the other 99 rows.
    """
    lower = filename.lower()
    if lower.endswith(".csv"):
        rows = _read_csv(content)
    elif lower.endswith((".xlsx", ".xlsm")):
        rows = _read_excel(content)
    else:
        raise ImportError_(
            "Unsupported file type. Upload a .csv or .xlsx file."
        )

    result = ParseResult()
    columns: dict[str, int] | None = None
    seen_in_file: set[str] = set()

    for index, raw_row in enumerate(rows, start=1):
        # Skip entirely blank rows — trailing ones are common in exports
        if not raw_row or all(cell in (None, "") for cell in raw_row):
            continue

        if columns is None:
            columns = _check_headers([_normalise_header(c) for c in raw_row])
            continue

        if len(result.leads) + len(result.rejected) >= MAX_ROWS:
            result.rejected.append(RejectedRow(
                row=index,
                reason=f"File exceeds the {MAX_ROWS} row limit — split it and upload again",
            ))
            break

        def cell(name: str) -> str:
            position = columns.get(name)
            if position is None or position >= len(raw_row):
                return ""
            return str(raw_row[position] or "").strip()

        name = cell("name")
        phone_raw = cell("phone")
        language = cell("language").lower() or DEFAULT_LANGUAGE

        if not name:
            result.rejected.append(RejectedRow(
                row=index, reason="Name is empty", raw=phone_raw,
            ))
            continue

        phone = normalise_phone(phone_raw)
        if not phone:
            result.rejected.append(RejectedRow(
                row=index,
                reason=f"'{phone_raw}' is not a valid Indian mobile number",
                raw=name,
            ))
            continue

        # Duplicated inside this same file
        key = phone_digits(phone)
        if key in seen_in_file:
            result.rejected.append(RejectedRow(
                row=index, reason="Duplicate of an earlier row in this file", raw=name,
            ))
            continue
        seen_in_file.add(key)

        if language not in VALID_LANGUAGES:
            # Not worth rejecting a row over — fall back and carry on
            language = DEFAULT_LANGUAGE

        result.leads.append(ParsedLead(
            row=index, name=name, phone=phone, language=language,
        ))

    if columns is None:
        raise ImportError_("The file appears to be empty")

    return result


def template_csv() -> str:
    """The template users download, with a couple of example rows."""
    return (
        "name,phone,language\n"
        "Arjun Mehta,+91 98765 43210,english\n"
        "Priya Sharma,9876543211,hindi\n"
        "Rohit Verma,+91 98765 43212,telugu\n"
    )
